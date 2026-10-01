#!/usr/bin/env python3
"""Generate test/fixtures/suite-v3.9.4.json from the AI Governance suite v3.9.4 sources.

Usage:  python3 scripts/generate-fixtures.py <folder holding the v3.9.4 .xlsx/.docx sources>

Reads (read-only):
  AIG-ASS-01_AGPI_Triage_Tool_Proposed.xlsx          sheet "AGPI Triage" A10:B15 (dimensions, weights), B17 thresholds,
                                                      B23 (§4.4.6 trigger answer) and the B17 trigger floor (v1.4)
  AIG-ASS-02_AI_Risk_Assessment_Worksheet_Proposed.xlsx sheet "Risk Assessment" A23:A27 (impact dimensions), C37:C43, C82 (agentic floor)
  AIG-AGT-02_Agentic_Classification_Reference.docx   Table A (per-dimension floors), Table B (rules), Table C (examples)
  AIG-AGT-03_Agentic_Triage_Function_Spec.docx       §6 agency tier -> minimum risk-tier pathway table

The expected grids are computed with a direct transcription of the workbook formulas
(ASS-01 D10:D16 and B17; ASS-02 C28, C37-C43 and E41). The transcription was checked
against LibreOffice recalculation of the real v3.9 workbooks on 54 sampled cases
(boundary scores, priority-floor cases, all evidence states, trigger and impact floors)
with 0 mismatches, and re-run on the v3.9.1 and v3.9.2 workbooks (AIG-ASS-02 v1.9 and v1.10) with the
same results; those cases (AIG-ASS-01 v1.4 with the row 23 trigger answer No / Yes / Unsure), plus a
504-case AIG-ASS-02 v1.10 agentic grid (C43, C82), are stored in test/fixtures/libreoffice-checked.json.
The v3.9.3 and v3.9.4 AIG-ASS-01 and AIG-ASS-02 workbooks are cell-for-cell identical to v3.9.2, so those cases still apply.
"""
import itertools, json, os, re, sys

import docx
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "."
OUT = os.path.join(os.path.dirname(__file__), "..", "test", "fixtures", "suite-v3.9.4.json")
TIERS = ["Low", "Medium", "High", "Critical"]
F01 = "AIG-ASS-01_AGPI_Triage_Tool_Proposed.xlsx"
F02 = "AIG-ASS-02_AI_Risk_Assessment_Worksheet_Proposed.xlsx"
F_AGT02 = "AIG-AGT-02_Agentic_Classification_Reference.docx"
F_AGT03 = "AIG-AGT-03_Agentic_Triage_Function_Spec.docx"


def ass01():
    wb = openpyxl.load_workbook(os.path.join(SRC, F01))
    ws = wb["AGPI Triage"]
    dims = [ws[f"A{r}"].value for r in range(10, 16)]
    weights = [round(ws[f"B{r}"].value * 100) for r in range(10, 16)]
    b17 = ws["B17"].value
    version = wb["Document Control"]["B7"].value
    assert 'IF(D16>=80,"Priority 1 – Critical",IF(OR(D16>=60,C10=5,C12=5),"Priority 2 – High"' in b17, "B17 changed"
    # v1.4 (W-05): row 23 trigger answer; a trigger use (Yes or Unsure) cannot be Priority 5.
    assert 'IF(OR(D16>=20,B23="Yes",B23="Unsure"),"Priority 4 – Routine","Priority 5 – Observe")' in b17, "B17 trigger floor changed"
    trigger_label = ws["A23"].value
    trigger_list = next(dv.formula1.strip('"').split(",") for dv in ws.data_validations.dataValidation if str(dv.sqref) == "B23")
    scores, prios, prios_trig = [], [], []
    for sc in itertools.product(range(1, 6), repeat=6):
        d16 = sum((c - 1) / 4 * w for c, w in zip(sc, weights))
        for trig, out in ((False, prios), (True, prios_trig)):
            if d16 >= 80: p = 1
            elif d16 >= 60 or sc[0] == 5 or sc[2] == 5: p = 2
            elif d16 >= 40: p = 3
            elif d16 >= 20 or trig: p = 4
            else: p = 5
            out.append(str(p))
        scores.append(round(d16 * 100))
    return {"source": f"{F01} (version {version}), sheet 'AGPI Triage', A10:B15, D10:D16, B17, B23",
            "dimensions": dims, "weights": weights,
            "trigger_input": {"label": trigger_label, "list": trigger_list},
            "grid_order": "itertools.product(1..5, repeat=6) over the six dimensions in workbook order",
            "score_x100": scores, "priority": "".join(prios), "priority_with_trigger": "".join(prios_trig)}


def ass02_tier(I, L, C, evidence, trig):
    inh = L * I
    inh_t = "Critical" if inh >= 16 else "High" if inh >= 11 else "Medium" if inh >= 6 else "Low"
    res = inh * (C / 5)
    res_t = "Critical" if res > 15 else "High" if res > 10 else "Medium" if res > 5 else "Low"
    elig = evidence == 2 or (evidence == 1 and inh_t not in ("High", "Critical"))
    floor = {"none": "Low", "special": "High", "vulnerable": "High", "housingCare": "High", "novel": "High",
             "materialChange": "High", "statutory": "Critical", "agenticNoReview": "Critical"}[trig]
    basis = res_t if elig else inh_t
    r = max(TIERS.index(basis), TIERS.index(floor), 1 if I == 5 else 0)
    return inh_t, res_t, TIERS[r]


def ass02():
    wb = openpyxl.load_workbook(os.path.join(SRC, F02))
    ws = wb["Risk Assessment"]
    dims = [ws[f"A{r}"].value for r in range(23, 28)]
    c43 = ws["C43"].value
    assert "IF(C28=5,2,1)" in c43 and 'IF(E41="Yes",C41,C38)' in c43, "C43 changed"
    assert ws["C41"].value.startswith('=IF(C40="","",IF(C40>15,"Critical",IF(C40>10,"High",IF(C40>5,"Medium","Low"))))')
    assert ws["C38"].value.startswith('=IF(C37="","",IF(C37>=16,"Critical",IF(C37>=11,"High",IF(C37>=6,"Medium","Low"))))')
    version = re.search(r"AIG-ASS-02 (v[\d.]+)", ws["A2"].value).group(1)
    cases = []
    # All seven §4.4.6 triggers (C50:C56); C42: statutory or agentic -> Critical, any other -> High.
    c42 = ws["C42"].value
    assert 'IF(OR(C54="Yes",C56="Yes",C56="Unsure"),"Critical",IF(COUNTIF(C50:C56,"Yes")>0,"High","Low"))' in c42, "C42 changed"
    triggers = [ws[f"A{r}"].value for r in range(50, 57)]
    for I, L, C, ev, trig in itertools.product(range(1, 6), range(1, 6), range(1, 6), (0, 1, 2),
                                               ("none", "special", "vulnerable", "housingCare", "novel", "statutory",
                                                "materialChange", "agenticNoReview")):
        inh_t, res_t, eff = ass02_tier(I, L, C, ev, trig)
        cases.append([I, L, C, ev, trig, inh_t, res_t, eff])
    # C82 agentic floor (v1.10, W-09): the agency tier token in C73 (T0-T5) is matched exactly and mapped
    # by INDEX; T4 is raised to Critical only by the C56 trigger floor (C42) where actions run without
    # evidenced per-action review. The pathway text in C74 is read only when C73 has no token.
    c82 = ws["C82"].value
    toks = re.search(r'MATCH\(LEFT\(C73,2\),\{([^}]+)\},0\)', c82).group(1).replace('"', "").split(",")
    vals = re.search(r'INDEX\(\{([^}]+)\},MATCH', c82).group(1).replace('"', "").split(",")
    assert 'SEARCH("Critical",C74)' not in c82, "C82 must not search free text for Critical"
    assert 'IF(C82="Critical",4,IF(C82="High",3,IF(C82="Medium",2,1)))' in c43, "C43 agentic term changed"
    agentic_floor = dict(zip(toks, vals))
    return {"source": f"{F02} ({version}), sheet 'Risk Assessment', A23:A27, C28, C37:C43, E41, A50:A56, C82",
            "agentic_floor_c82": agentic_floor,
            "triggers": triggers,
            "impact_dimensions": dims,
            "evidence_codes": {"0": "E37 No", "1": "E37 Yes + ref, E39 No (no independent check)",
                               "2": "E37 Yes + ref, E39 Yes + ref (independently verified)"},
            "columns": ["maxImpact", "likelihood", "control", "evidence", "trigger",
                        "inherentTier", "residualTier", "effectiveTier"],
            "cases": cases}


def cell_rows(t):
    return [[c.text.strip() for c in r.cells] for r in t.rows]


def agt02():
    d = docx.Document(os.path.join(SRC, F_AGT02))
    ta = next(t for t in d.tables if [c.text for c in t.rows[0].cells][:2] == ["Score", "Consequence"])
    rows = cell_rows(ta)
    head = [h.lower() for h in rows[0][1:]]
    table_a = {h: [int(rows[i][j + 1][1]) for i in range(1, 7)] for j, h in enumerate(head)}
    tb = next(t for t in d.tables if [c.text for c in t.rows[0].cells] == ["Rule", "Condition", "Tier floor"])
    table_b = {r[0]: {"condition": r[1], "floor": int(r[2][1])} for r in cell_rows(tb)[1:]}
    tc = next(t for t in d.tables if t.rows[0].cells[0].text == "Example")
    examples = []
    for r in cell_rows(tc)[1:]:
        m = re.match(r"(\d) / (\d) / (\d) / (\d) / (\d); flags: (.*)", r[1])
        flags = [] if m.group(6).strip() == "none" else [f.strip() for f in m.group(6).split(",")]
        tier = int(re.search(r"→ T(\d)", r[2]).group(1))
        examples.append({"example": r[0], "scores": [int(m.group(i)) for i in range(1, 6)], "flags": flags,
                         "rules_listed": re.findall(r"rules? ([A-D]\d(?:, [A-D]\d)*)", r[2]),
                         "agency_tier": tier, "agency_minimum": r[3], "assumed_risk": r[4], "governing_tier": r[5]})
    return {"source": f"{F_AGT02}, Tables A, B and C", "table_a": table_a, "table_b": table_b,
            "calibration_examples": examples}


def agt03():
    d = docx.Document(os.path.join(SRC, F_AGT03))
    t = next(t for t in d.tables if t.rows[0].cells[0].text == "Agency tier")
    rows = cell_rows(t)[1:]
    return {"source": f"{F_AGT03}, §6 pathway table",
            "minimum_pathway": {r[0].split()[0]: r[1] for r in rows},
            "agentic_requirement": {r[0].split()[0]: r[2] for r in rows}}


if __name__ == "__main__":
    out = {"suite": "v3.9.4 (1 October 2026)", "ass01": ass01(), "ass02": ass02(), "agt02": agt02(), "agt03": agt03()}
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print("wrote", os.path.normpath(OUT))
